"""Diller sanovi — tayyor ro'yxat (ведомость) oqimi.

prefill (manbalar, snapshot, takror), claim qulfi, sanalganlarni idempotent yozish,
yuborishda sanalmaganlar (zero/keep), solishtiruvda "sanalmagan" bayrog'i, ledger
o'zgarmasligi. Smartup so'rovi monkeypatch — tarmoqqa chiqilmaydi.
"""
from __future__ import annotations

import uuid

from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.models.order import Order, OrderLine, OrderWmsState
from app.models.product import Product
from app.models.settings_organization import SettingsOrganization
from app.models.stock import StockMovement
from app.services import dealer_count_compare as cmp
from app.services import dealer_count_sheet as sheet
from tests.test_sale_expiry_cutoff import _as, _clear, _mk_user

URL = "/api/v1/dealer-counts"


def _product(db: Session, tag: str) -> Product:
    p = Product(
        external_source="t",
        external_id=f"{tag}-{uuid.uuid4().hex[:6]}",
        name=f"Prod {tag}",
        sku=f"{tag}-{uuid.uuid4().hex[:5]}",
        barcode=f"46{uuid.uuid4().int % 10**11:011d}",
        is_active=True,
    )
    db.add(p)
    db.commit()
    db.refresh(p)
    return p


def _org(db: Session, *, wh: str | None) -> SettingsOrganization:
    org = SettingsOrganization(org_id=f"30{uuid.uuid4().hex[:6]}", name="Дилер Sheet", smartup_warehouse_code=wh)
    db.add(org)
    db.commit()
    db.refresh(org)
    return org


def _new_sheet(client: TestClient, org) -> str:
    res = client.post(URL, json={"client_uuid": str(uuid.uuid4()), "dealer_org_id": org.org_id, "source": "web", "lines": []})
    assert res.status_code == 200, res.text
    return res.json()["id"]


def _patch_smartup(monkeypatch, balance: dict[str, float]):
    monkeypatch.setattr(sheet, "load_dealer_balance", lambda db, org_id, refresh=False: (dict(balance), "cache", "wh30"))
    monkeypatch.setattr(cmp, "load_dealer_balance", lambda db, org_id, refresh=False: (dict(balance), "cache", "wh30"))


# --- prefill -----------------------------------------------------------------


def test_prefill_smartup_adds_snapshot_and_skips_duplicates(client: TestClient, db_session: Session, monkeypatch):
    org = _org(db_session, wh="wh30")
    a, b = _product(db_session, "A"), _product(db_session, "B")
    _patch_smartup(monkeypatch, {a.sku: 10, b.sku: 5, "NOT-IN-CATALOG": 3})
    _as(_mk_user(db_session, "warehouse_admin"))
    try:
        cid = _new_sheet(client, org)
        r1 = client.post(f"{URL}/{cid}/prefill", json={"sources": ["smartup"]})
        assert r1.status_code == 200, r1.text
        assert (r1.json()["added"], r1.json()["skipped"], r1.json()["not_in_catalog"]) == (2, 0, 1)
        c = r1.json()["count"]
        assert c["source"] == "sheet"
        assert c["sheet_lines"] == 2 and c["counted_lines"] == 0 and c["lines_count"] == 0
        by = {ln["sku"]: ln for ln in c["lines"]}
        assert by[a.sku]["qty"] is None and float(by[a.sku]["snapshot_qty"]) == 10
        assert by[b.sku]["counted_at"] is None

        r2 = client.post(f"{URL}/{cid}/prefill", json={"sources": ["smartup"]})
        assert (r2.json()["added"], r2.json()["skipped"]) == (0, 2)
    finally:
        _clear()


def test_prefill_falls_back_to_shipped_without_warehouse_code(client: TestClient, db_session: Session):
    org = _org(db_session, wh=None)
    x = _product(db_session, "X")
    order = Order(source="diller", source_external_id=f"d-{uuid.uuid4().hex[:8]}", order_number=f"D-{uuid.uuid4().hex[:6]}", to_filial_code=org.org_id)
    order.wms_state = OrderWmsState(status="completed")
    order.lines = [OrderLine(sku=x.sku, name="X line", qty=3.0, uom="dona")]
    db_session.add(order)
    db_session.commit()
    _as(_mk_user(db_session, "warehouse_admin"))
    try:
        cid = _new_sheet(client, org)
        r = client.post(f"{URL}/{cid}/prefill", json={"sources": ["smartup"]})
        assert r.status_code == 200, r.text
        assert r.json()["sources"] == ["shipped"]
        assert r.json()["added"] == 1
        ln = r.json()["count"]["lines"][0]
        assert ln["sku"] == x.sku and ln["qty"] is None and ln["snapshot_qty"] is None
    finally:
        _clear()


# --- claim / lock ---------------------------------------------------------------


def test_claim_locks_web_edit_and_other_users(client: TestClient, db_session: Session, monkeypatch):
    org = _org(db_session, wh="wh30")
    a = _product(db_session, "A")
    _patch_smartup(monkeypatch, {a.sku: 4})
    admin = _mk_user(db_session, "warehouse_admin")
    u1 = _mk_user(db_session, "inventory_controller")
    u2 = _mk_user(db_session, "inventory_controller")
    _as(admin)
    try:
        cid = _new_sheet(client, org)
        client.post(f"{URL}/{cid}/prefill", json={"sources": ["smartup"]})
    finally:
        _clear()
    _as(u1)
    try:
        r = client.post(f"{URL}/{cid}/claim")
        assert r.status_code == 200 and r.json()["status"] == "in_progress"
        assert r.json()["assigned_to_user_id"] == str(u1.id)
        assert client.post(f"{URL}/{cid}/claim").status_code == 200  # o'zi qayta olsa bo'ladi
    finally:
        _clear()
    _as(u2)
    try:
        assert client.post(f"{URL}/{cid}/claim").status_code == 409
        assert client.put(f"{URL}/{cid}/counts", json={"entries": [{"product_id": str(a.id), "qty": 1}]}).status_code == 403
    finally:
        _clear()
    _as(admin)
    try:
        assert client.put(f"{URL}/{cid}", json={"lines": []}).status_code == 409  # web qulf
        assert "telefonda" in client.put(f"{URL}/{cid}", json={"lines": []}).text
        rel = client.post(f"{URL}/{cid}/release")
        assert rel.status_code == 200 and rel.json()["status"] == "draft" and rel.json()["assigned_to_user_id"] is None
    finally:
        _clear()


# --- counts (idempotent) ------------------------------------------------------


def test_counts_are_idempotent_and_can_add_new_products(client: TestClient, db_session: Session, monkeypatch):
    org = _org(db_session, wh="wh30")
    a, b, extra = _product(db_session, "A"), _product(db_session, "B"), _product(db_session, "E")
    _patch_smartup(monkeypatch, {a.sku: 10, b.sku: 5})
    u = _mk_user(db_session, "inventory_controller")
    _as(_mk_user(db_session, "warehouse_admin"))
    try:
        cid = _new_sheet(client, org)
        client.post(f"{URL}/{cid}/prefill", json={"sources": ["smartup"]})
    finally:
        _clear()
    _as(u)
    try:
        client.post(f"{URL}/{cid}/claim")
        entries = [{"product_id": str(a.id), "qty": 7}, {"product_id": str(extra.id), "qty": 2}]
        r1 = client.put(f"{URL}/{cid}/counts", json={"entries": entries})
        assert r1.status_code == 200, r1.text
        assert (r1.json()["updated"], r1.json()["added"]) == (1, 1)
        r2 = client.put(f"{URL}/{cid}/counts", json={"entries": entries})
        assert (r2.json()["updated"], r2.json()["added"]) == (2, 0)
        c = r2.json()["count"]
        by = {ln["sku"]: ln for ln in c["lines"]}
        assert float(by[a.sku]["qty"]) == 7 and by[a.sku]["counted_at"] is not None  # qo'shilmagan, almashgan
        assert float(by[extra.sku]["qty"]) == 2
        assert by[b.sku]["qty"] is None
        assert c["sheet_lines"] == 3 and c["counted_lines"] == 2 and c["lines_count"] == 2
    finally:
        _clear()


# --- submit: uncounted zero / keep -----------------------------------------------


def _sheet_with_one_counted(client, db_session, monkeypatch, counter):
    """Ro'yxatni web (admin) tayyorlaydi, `counter` telefonda oladi va bitta qatorni sanaydi.
    Qaytganda `counter` nomidan so'rov yuborilmoqda (chaqiruvchi `_clear()` qiladi)."""
    org = _org(db_session, wh="wh30")
    a, b, c = _product(db_session, "A"), _product(db_session, "B"), _product(db_session, "C")
    _patch_smartup(monkeypatch, {a.sku: 10, b.sku: 5, c.sku: 2})
    _as(_mk_user(db_session, "warehouse_admin"))
    try:
        cid = _new_sheet(client, org)
        client.post(f"{URL}/{cid}/prefill", json={"sources": ["smartup"]})
    finally:
        _clear()
    _as(counter)
    client.post(f"{URL}/{cid}/claim")
    client.put(f"{URL}/{cid}/counts", json={"entries": [{"product_id": str(a.id), "qty": 8}]})
    return cid, a, b, c


def test_submit_uncounted_zero_keeps_counted_at_empty(client: TestClient, db_session: Session, monkeypatch):
    try:
        cid, a, b, c = _sheet_with_one_counted(client, db_session, monkeypatch, _mk_user(db_session, "inventory_controller"))
        r = client.post(f"{URL}/{cid}/submit", params={"uncounted": "zero"})
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["status"] == "submitted" and body["uncounted_policy"] == "zero"
        by = {ln["sku"]: ln for ln in body["lines"]}
        assert float(by[b.sku]["qty"]) == 0 and by[b.sku]["counted_at"] is None
        assert float(by[a.sku]["qty"]) == 8 and by[a.sku]["counted_at"] is not None
        assert body["lines_count"] == 3 and body["counted_lines"] == 1 and float(body["total_units"]) == 8

        cmpr = client.get(f"{URL}/{cid}/compare").json()
        rows = {r["sku"]: r for r in cmpr["rows"]}
        assert rows[a.sku]["is_counted"] is True and rows[a.sku]["diff"] == -2
        assert rows[b.sku]["is_counted"] is False and rows[b.sku]["counted"] == 0 and rows[b.sku]["diff"] == -5
        assert cmpr["uncounted_skus"] == 2
    finally:
        _clear()


def test_submit_uncounted_keep_leaves_null(client: TestClient, db_session: Session, monkeypatch):
    try:
        cid, a, b, _c = _sheet_with_one_counted(client, db_session, monkeypatch, _mk_user(db_session, "inventory_controller"))
        r = client.post(f"{URL}/{cid}/submit", params={"uncounted": "keep"})
        assert r.status_code == 200, r.text
        by = {ln["sku"]: ln for ln in r.json()["lines"]}
        assert by[b.sku]["qty"] is None
        assert r.json()["lines_count"] == 1
        assert client.post(f"{URL}/{cid}/submit", params={"uncounted": "bad"}).status_code == 409  # allaqachon yuborilgan
    finally:
        _clear()


def test_submit_rejects_sheet_with_nothing_counted(client: TestClient, db_session: Session, monkeypatch):
    org = _org(db_session, wh="wh30")
    a = _product(db_session, "A")
    _patch_smartup(monkeypatch, {a.sku: 10})
    _as(_mk_user(db_session, "warehouse_admin"))
    try:
        cid = _new_sheet(client, org)
        client.post(f"{URL}/{cid}/prefill", json={"sources": ["smartup"]})
        assert client.post(f"{URL}/{cid}/submit", params={"uncounted": "keep"}).status_code == 400
    finally:
        _clear()


def test_sheet_flow_never_touches_stock_ledger(client: TestClient, db_session: Session, monkeypatch):
    before = db_session.query(StockMovement).count()
    try:
        cid, _a, _b, _c = _sheet_with_one_counted(client, db_session, monkeypatch, _mk_user(db_session, "warehouse_admin"))
        assert client.post(f"{URL}/{cid}/submit").status_code == 200
    finally:
        _clear()
    assert db_session.query(StockMovement).count() == before


def test_mobile_created_count_marks_lines_counted(client: TestClient, db_session: Session):
    org = _org(db_session, wh=None)
    a = _product(db_session, "A")
    _as(_mk_user(db_session, "warehouse_admin"))
    try:
        r = client.post(URL, json={"client_uuid": str(uuid.uuid4()), "dealer_org_id": org.org_id, "submit": True,
                                   "lines": [{"product_id": str(a.id), "qty": 3}]})
        assert r.status_code == 200, r.text
        assert r.json()["source"] == "mobile" and r.json()["counted_lines"] == 1
        assert r.json()["lines"][0]["counted_at"] is not None
    finally:
        _clear()


def test_counter_cannot_create_or_prefill_sheet(client: TestClient, db_session: Session, monkeypatch):
    """Ro'yxat faqat web'da: telefondagi sanovchi yaratolmaydi va to'ldirolmaydi."""
    org = _org(db_session, wh="wh30")
    a = _product(db_session, "A")
    _patch_smartup(monkeypatch, {a.sku: 1})
    _as(_mk_user(db_session, "warehouse_admin"))
    try:
        cid = _new_sheet(client, org)
    finally:
        _clear()
    _as(_mk_user(db_session, "inventory_controller"))
    try:
        for body in (
            {"source": "sheet", "lines": []},
            {"source": "mobile", "lines": []},
            # submit bilan ham, lekin mobil draft emas — web yo'li.
            {"source": "web", "submit": True, "lines": [{"product_id": str(a.id), "qty": 1}]},
        ):
            r = client.post(URL, json={"client_uuid": str(uuid.uuid4()), "dealer_org_id": org.org_id, **body})
            assert r.status_code == 403, body
            assert "faqat web" in r.text
        assert client.post(f"{URL}/{cid}/prefill", json={"sources": ["smartup"]}).status_code == 403
    finally:
        _clear()


def test_mine_includes_sheets_counted_by_me_and_export_names_counter(
    client: TestClient, db_session: Session, monkeypatch
):
    counter = _mk_user(db_session, "inventory_controller")
    try:
        cid, _a, _b, _c = _sheet_with_one_counted(client, db_session, monkeypatch, counter)
        assert client.post(f"{URL}/{cid}/submit", params={"uncounted": "keep"}).status_code == 200
        mine = client.get(URL, params={"mine": True}).json()
        assert cid in {it["id"] for it in mine["items"]}

        res = client.get(f"{URL}/{cid}/export.xlsx")  # sanalmagan (qty yo'q) qatorlar bilan
        assert res.status_code == 200, res.text
        from io import BytesIO

        from openpyxl import load_workbook

        rows = {r[0]: r[1] for r in load_workbook(BytesIO(res.content)).active.iter_rows(values_only=True) if r}
        assert rows["Sanadi"] == (counter.full_name or counter.username)
    finally:
        _clear()


def test_only_admin_deletes_sheet_taken_by_phone(client: TestClient, db_session: Session, monkeypatch):
    counter = _mk_user(db_session, "inventory_controller")
    try:
        cid, _a, _b, _c = _sheet_with_one_counted(client, db_session, monkeypatch, counter)
        # Sanovchi o'zi olgan ro'yxatni ham o'chirolmaydi (web qulf qoidasi).
        assert client.delete(f"{URL}/{cid}").status_code == 409
    finally:
        _clear()
    _as(_mk_user(db_session, "warehouse_admin"))
    try:
        assert client.delete(f"{URL}/{cid}").status_code == 204
    finally:
        _clear()
    _as(counter)
    try:
        # Telefon keyingi so'rovda 404 oladi — nusxasini o'chirish belgisi.
        assert client.put(f"{URL}/{cid}/counts", json={"entries": []}).status_code == 404
        assert client.post(f"{URL}/{cid}/submit").status_code == 404
    finally:
        _clear()


def test_available_hides_sheets_taken_by_others(client: TestClient, db_session: Session, monkeypatch):
    """Telefon ro'yxati: ochiq + men olgan; boshqa xodim olgani ko'rinmaydi."""
    a = _product(db_session, "A")
    _patch_smartup(monkeypatch, {a.sku: 1})
    org_free, org_mine, org_other = _org(db_session, wh="wh30"), _org(db_session, wh="wh30"), _org(db_session, wh="wh30")
    me, other = _mk_user(db_session, "inventory_controller"), _mk_user(db_session, "inventory_controller")
    ids = {}
    _as(_mk_user(db_session, "warehouse_admin"))
    try:
        for key, org in (("free", org_free), ("mine", org_mine), ("other", org_other)):
            ids[key] = _new_sheet(client, org)
            client.post(f"{URL}/{ids[key]}/prefill", json={"sources": ["smartup"]})
    finally:
        _clear()
    _as(other)
    try:
        assert client.post(f"{URL}/{ids['other']}/claim").status_code == 200
    finally:
        _clear()
    _as(me)
    try:
        assert client.post(f"{URL}/{ids['mine']}/claim").status_code == 200
        got = {it["id"] for it in client.get(URL, params={"status": "draft,in_progress", "available": True, "limit": 500}).json()["items"]}
        assert ids["free"] in got and ids["mine"] in got
        assert ids["other"] not in got
    finally:
        _clear()


def test_list_accepts_multiple_statuses(client: TestClient, db_session: Session, monkeypatch):
    org = _org(db_session, wh="wh30")
    a = _product(db_session, "A")
    _patch_smartup(monkeypatch, {a.sku: 1})
    _as(_mk_user(db_session, "warehouse_admin"))
    try:
        cid = _new_sheet(client, org)
        client.post(f"{URL}/{cid}/prefill", json={"sources": ["smartup"]})
        client.post(f"{URL}/{cid}/claim")
        r = client.get(URL, params={"status": "draft,in_progress", "dealer_org_id": org.org_id})
        assert r.json()["total"] == 1 and r.json()["items"][0]["status"] == "in_progress"
        assert client.get(URL, params={"status": "submitted", "dealer_org_id": org.org_id}).json()["total"] == 0
    finally:
        _clear()
