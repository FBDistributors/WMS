"""Diller ombor qoldig'i sanovi: WMS ledgeridan alohida hujjatlar.

Kafolatlar: yaratish faqat web'da va idempotent, dillerda bitta faol sanov (yangisi eskisini
yopadi), qator darajasidagi web tahriri, skan serverda resolve bo'ladi, o'chirish huquqlari,
eski ilova qadamlari "ilovani yangilang" beradi va sanov `stock_movements` ga hech narsa yozmaydi.
"""
from __future__ import annotations

import uuid

from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.models.product import Product
from app.models.settings_organization import SettingsOrganization
from app.models.stock import StockMovement
from tests.test_sale_expiry_cutoff import _as, _clear, _mk_user

URL = "/api/v1/dealer-counts"


def _seed(db: Session):
    org = SettingsOrganization(org_id=f"81{uuid.uuid4().hex[:6]}", name="Дилер Test (Sinov)")
    head = SettingsOrganization(org_id="3788131", name="Головной Офис")
    product = Product(
        external_source="test",
        external_id=f"dc-{uuid.uuid4().hex[:8]}",
        name="Dealer Count Prod",
        sku=f"SKU-DC-{uuid.uuid4().hex[:6]}",
        barcode=f"46{uuid.uuid4().int % 10**11:011d}",
        is_active=True,
    )
    db.add_all([org, head, product])
    db.commit()
    db.refresh(org)
    db.refresh(product)
    return org, product


def _payload(org, product=None, *, client_uuid=None, lines=None, replace=False):
    return {
        "client_uuid": str(client_uuid or uuid.uuid4()),
        "dealer_org_id": org.org_id,
        "replace": replace,
        "lines": lines
        if lines is not None
        else ([{"product_id": str(product.id), "scanned_barcode": product.barcode}] if product else []),
    }


def test_dealers_list_hides_head_office(client: TestClient, db_session: Session):
    org, _ = _seed(db_session)
    _as(_mk_user(db_session, "inventory_controller"))
    try:
        ids = {d["org_id"] for d in client.get(f"{URL}/dealers").json()}
        assert org.org_id in ids
        assert "3788131" not in ids
    finally:
        _clear()


def test_create_is_idempotent_and_active(client: TestClient, db_session: Session):
    org, product = _seed(db_session)
    _as(_mk_user(db_session, "warehouse_admin"))
    try:
        cu = uuid.uuid4()
        a = client.post(URL, json=_payload(org, product, client_uuid=cu))
        b = client.post(URL, json=_payload(org, product, client_uuid=cu))
        assert a.status_code == 200, a.text
        assert b.status_code == 200 and a.json()["id"] == b.json()["id"]
        body = a.json()
        assert body["is_active"] is True
        assert body["dealer_name"] == org.name
        assert body["sheet_lines"] == 1 and body["counted_lines"] == 0
        assert "status" not in body
        assert client.get(URL, params={"dealer_org_id": org.org_id}).json()["total"] == 1
    finally:
        _clear()


def test_counter_cannot_create(client: TestClient, db_session: Session):
    org, product = _seed(db_session)
    for role in ("inventory_controller", "picker"):
        _as(_mk_user(db_session, role))
        try:
            assert client.post(URL, json=_payload(org, product)).status_code == 403
        finally:
            _clear()


def test_new_count_replaces_active_one(client: TestClient, db_session: Session):
    """Dillerda bitta faol sanov: tasdiqsiz 409, `replace` bilan eskisi yopiladi (o'chmaydi)."""
    org, product = _seed(db_session)
    _as(_mk_user(db_session, "warehouse_admin"))
    try:
        first = client.post(URL, json=_payload(org, product)).json()
        again = client.post(URL, json=_payload(org, product))
        assert again.status_code == 409 and "faol sanov" in again.text
        second = client.post(URL, json=_payload(org, product, replace=True))
        assert second.status_code == 200, second.text
        assert client.get(f"{URL}/{first['id']}").json()["is_active"] is False
        active = client.get(URL, params={"active": True, "dealer_org_id": org.org_id}).json()["items"]
        assert [c["id"] for c in active] == [second.json()["id"]]
    finally:
        _clear()


def test_web_line_edits_keep_phone_counts(client: TestClient, db_session: Session):
    """Web qator qo'shsa / boshqa qatorni o'zgartirsa — telefon sanagan qator o'z joyida qoladi."""
    org, product = _seed(db_session)
    other = Product(external_source="t", external_id=f"o-{uuid.uuid4().hex[:6]}", name="Other",
                    sku=f"SKU-O-{uuid.uuid4().hex[:5]}", barcode="4600000000017", is_active=True)
    db_session.add(other)
    db_session.commit()
    admin, counter = _mk_user(db_session, "warehouse_admin"), _mk_user(db_session, "inventory_controller")
    _as(admin)
    try:
        cid = client.post(URL, json=_payload(org, product)).json()["id"]
    finally:
        _clear()
    _as(counter)
    try:
        assert client.put(f"{URL}/{cid}/counts", json={"entries": [{"product_id": str(product.id), "qty": 7}]}).status_code == 200
    finally:
        _clear()
    _as(admin)
    try:
        r = client.post(f"{URL}/{cid}/lines", json={"lines": [{"product_id": str(other.id), "scanned_barcode": ""}]})
        assert r.status_code == 200, r.text
        lines = {ln["sku"]: ln for ln in r.json()["count"]["lines"]}
        assert float(lines[product.sku]["qty"]) == 7
        assert lines[product.sku]["counted_by_name"]
        # Qatorni o'zgartirish / sanalmaganga qaytarish / o'chirish.
        oid = lines[other.sku]["id"]
        p = client.patch(f"{URL}/{cid}/lines/{oid}", json={"qty": 3, "location_code": " a-3 "})
        assert p.status_code == 200, p.text
        ln = next(x for x in p.json()["lines"] if x["id"] == oid)
        assert float(ln["qty"]) == 3 and ln["location_code"] == "A-3" and ln["counted_at"]
        ln = next(x for x in client.patch(f"{URL}/{cid}/lines/{oid}", json={"qty": None}).json()["lines"] if x["id"] == oid)
        assert ln["qty"] is None and ln["counted_at"] is None and ln["location_code"] == "A-3"
        d = client.delete(f"{URL}/{cid}/lines/{oid}")
        assert d.status_code == 200 and d.json()["sheet_lines"] == 1
    finally:
        _clear()
    _as(counter)
    try:
        # Telefondagi xodim ro'yxatni tahrirlay olmaydi.
        assert client.post(f"{URL}/{cid}/lines", json={"lines": []}).status_code == 403
    finally:
        _clear()


def test_barcode_resolved_server_side(client: TestClient, db_session: Session):
    org, product = _seed(db_session)
    _as(_mk_user(db_session, "warehouse_admin"))
    try:
        res = client.post(
            URL,
            json=_payload(
                org,
                lines=[
                    {"scanned_barcode": product.barcode, "qty": 5, "expiry_date": "2027-03-15"},
                    {"scanned_barcode": "0000000000000", "qty": 1},
                ],
            ),
        )
        assert res.status_code == 200, res.text
        by_sku = {ln["sku"]: ln for ln in res.json()["lines"]}
        assert by_sku[product.sku]["expiry_date"] == "2027-03-01"  # oy boshiga normallashadi
        assert by_sku[None]["product_id"] is None and by_sku[None]["scanned_barcode"] == "0000000000000"
    finally:
        _clear()


def test_count_never_touches_stock_ledger(client: TestClient, db_session: Session):
    org, product = _seed(db_session)
    _as(_mk_user(db_session, "warehouse_admin"))
    try:
        before = db_session.query(StockMovement).count()
        cid = client.post(URL, json=_payload(org, product)).json()["id"]
        client.put(f"{URL}/{cid}/counts", json={"entries": [{"product_id": str(product.id), "qty": 4}]})
        assert db_session.query(StockMovement).count() == before
    finally:
        _clear()


def test_delete_by_creator_or_admin(client: TestClient, db_session: Session):
    org, product = _seed(db_session)
    owner = _mk_user(db_session, "supervisor")
    other = _mk_user(db_session, "inventory_controller")
    admin = _mk_user(db_session, "warehouse_admin")
    _as(owner)
    try:
        cid = client.post(URL, json=_payload(org, product)).json()["id"]
    finally:
        _clear()
    _as(other)
    try:
        assert client.delete(f"{URL}/{cid}").status_code == 403
    finally:
        _clear()
    _as(admin)
    try:
        assert client.delete(f"{URL}/{cid}").status_code == 204
        assert client.get(f"{URL}/{cid}").status_code == 404
    finally:
        _clear()


def test_old_app_steps_ask_to_update(client: TestClient, db_session: Session):
    org, product = _seed(db_session)
    _as(_mk_user(db_session, "warehouse_admin"))
    try:
        cid = client.post(URL, json=_payload(org, product)).json()["id"]
        for step in ("claim", "release", "submit"):
            r = client.post(f"{URL}/{cid}/{step}")
            assert r.status_code == 410 and "yangilang" in r.text
    finally:
        _clear()


def test_export_xlsx_has_location_and_counter(client: TestClient, db_session: Session):
    org, product = _seed(db_session)
    _as(_mk_user(db_session, "warehouse_admin"))
    try:
        cid = client.post(URL, json=_payload(org, product)).json()["id"]
        client.put(f"{URL}/{cid}/counts", json={"entries": [{"product_id": str(product.id), "qty": 12, "location_code": "B-1"}]})
        res = client.get(f"{URL}/{cid}/export.xlsx")
        assert res.status_code == 200, res.text
        from io import BytesIO

        from openpyxl import load_workbook

        rows = list(load_workbook(BytesIO(res.content)).active.iter_rows(values_only=True))
        header_idx = next(i for i, r in enumerate(rows) if r and r[0] == "#")
        assert rows[header_idx][4] == "Joy"
        line = rows[header_idx + 1]
        assert line[1] == product.sku and line[4] == "B-1" and line[5] == 12 and line[7]
    finally:
        _clear()
